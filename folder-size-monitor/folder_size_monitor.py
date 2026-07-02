#!/usr/bin/env python3
"""Scan disk folders and export size report to Excel."""

from __future__ import annotations

import argparse
import os
import sys
import time
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


@dataclass
class FolderInfo:
    path: Path
    size: int
    parent: Path | None
    depth: int
    name: str
    inclusion_chain: str


def format_size(size_bytes: int) -> str:
    units = ("B", "KB", "MB", "GB", "TB", "PB")
    size = float(size_bytes)
    for unit in units:
        if size < 1024 or unit == units[-1]:
            if unit == "B":
                return f"{int(size)} {unit}"
            return f"{size:.2f} {unit}"
        size /= 1024
    return f"{size_bytes} B"


def scan_folder_sizes(root: Path) -> dict[Path, int]:
    """Return total size (bytes) for each folder under root, including subfolders."""
    direct_sizes: dict[Path, int] = {}
    root = root.resolve()

    for current_root, dirnames, filenames in os.walk(root, followlinks=False):
        current_path = Path(current_root)
        total = 0

        for filename in filenames:
            file_path = current_path / filename
            try:
                total += file_path.stat().st_size
            except (OSError, PermissionError):
                continue

        direct_sizes[current_path] = total

        # Skip inaccessible subdirectories early.
        accessible_dirs: list[str] = []
        for dirname in dirnames:
            child = current_path / dirname
            try:
                child.stat()
                accessible_dirs.append(dirname)
            except (OSError, PermissionError):
                direct_sizes[child] = 0
        dirnames[:] = accessible_dirs

    # Aggregate child folder sizes into parents (deepest paths first).
    all_paths = sorted(direct_sizes, key=lambda p: len(p.parts), reverse=True)
    totals = dict(direct_sizes)

    for path in all_paths:
        parent = path.parent
        if parent != path and parent in totals:
            totals[parent] += totals[path]

    return totals


def build_inclusion_chain(path: Path, root: Path) -> str:
    try:
        relative = path.resolve().relative_to(root.resolve())
        parts = [root.name or str(root)] + list(relative.parts)
    except ValueError:
        parts = list(path.parts)
    return " > ".join(parts)


def collect_folder_info(root: Path, sizes: dict[Path, int]) -> list[FolderInfo]:
    root = root.resolve()
    folders: list[FolderInfo] = []

    for path, size in sizes.items():
        parent = path.parent if path != root else None
        if parent is not None and parent not in sizes and parent != path:
            # Keep parent reference for display even if parent wasn't scanned as a key.
            pass

        try:
            depth = len(path.resolve().relative_to(root).parts)
        except ValueError:
            depth = len(path.parts)

        folders.append(
            FolderInfo(
                path=path,
                size=size,
                parent=parent,
                depth=depth,
                name=path.name or str(path),
                inclusion_chain=build_inclusion_chain(path, root),
            )
        )

    folders.sort(key=lambda item: item.size, reverse=True)
    return folders


def export_to_excel(folders: list[FolderInfo], output_path: Path, scan_root: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "文件夹大小"

    headers = [
        "排名",
        "路径",
        "文件夹名称",
        "大小(字节)",
        "大小",
        "父目录",
        "层级",
        "包含关系",
    ]

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for col, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for index, folder in enumerate(folders, start=1):
        parent_text = str(folder.parent) if folder.parent is not None else ""
        row = index + 1
        values = [
            index,
            str(folder.path),
            folder.name,
            folder.size,
            format_size(folder.size),
            parent_text,
            folder.depth,
            folder.inclusion_chain,
        ]
        for col, value in enumerate(values, start=1):
            sheet.cell(row=row, column=col, value=value)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:H{len(folders) + 1}"

    column_widths = {
        1: 8,
        2: 60,
        3: 28,
        4: 14,
        5: 14,
        6: 60,
        7: 8,
        8: 80,
    }
    for col, width in column_widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width

    summary = workbook.create_sheet("扫描信息")
    summary["A1"] = "扫描根目录"
    summary["B1"] = str(scan_root.resolve())
    summary["A2"] = "文件夹总数"
    summary["B2"] = len(folders)
    summary["A3"] = "总占用空间"
    summary["B3"] = format_size(sum(folder.size for folder in folders if folder.path == scan_root.resolve()))
    summary["A4"] = "生成时间"
    summary["B4"] = time.strftime("%Y-%m-%d %H:%M:%S")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def run_scan(scan_root: Path, output_path: Path) -> tuple[int, float]:
    """Scan folders and write Excel report. Returns (folder_count, elapsed_seconds)."""
    scan_root = scan_root.expanduser().resolve()
    output_path = output_path.expanduser()

    if not scan_root.exists():
        raise FileNotFoundError(f"路径不存在: {scan_root}")
    if not scan_root.is_dir():
        raise NotADirectoryError(f"路径不是文件夹: {scan_root}")

    start = time.time()
    sizes = scan_folder_sizes(scan_root)
    folders = collect_folder_info(scan_root, sizes)
    export_to_excel(folders, output_path, scan_root)
    elapsed = time.time() - start
    return len(folders), elapsed


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="扫描硬盘文件夹大小，并导出包含层级关系的 Excel 报表。",
    )
    parser.add_argument(
        "path",
        nargs="?",
        default=".",
        help="要扫描的目录路径（默认：当前目录）",
    )
    parser.add_argument(
        "-o",
        "--output",
        default="folder_sizes.xlsx",
        help="输出 Excel 文件路径（默认：folder_sizes.xlsx）",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    scan_root = Path(args.path)
    output_path = Path(args.output)

    try:
        print(f"正在扫描: {scan_root.expanduser().resolve()}")
        folder_count, elapsed = run_scan(scan_root, output_path)
    except (FileNotFoundError, NotADirectoryError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 1

    print(f"扫描完成，共 {folder_count} 个文件夹")
    print(f"结果已保存: {output_path.expanduser().resolve()}")
    print(f"耗时: {elapsed:.1f} 秒")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
